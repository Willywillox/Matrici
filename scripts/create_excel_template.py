#!/usr/bin/env python3
"""
Crea template Excel per import operatori
"""

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_template(output_path='templates/template_import_operatori.xlsx'):
    """Crea template Excel"""

    wb = Workbook()
    ws = wb.active
    ws.title = "Operatori"

    # Definisci colonne
    columns = [
        # === OBBLIGATORIE (sfondo rosso) ===
        ('Nome', 15, True),
        ('Cognome', 15, True),
        ('ID_SAP', 12, True),
        ('Data_Riferimento', 18, True),

        # === ANAGRAFICA (sfondo azzurro) ===
        ('Tipo_Contratto', 15, False),
        ('FTE', 8, False),
        ('Ore_Settimana', 15, False),

        # === TURNO (sfondo verde) ===
        ('ID_Turno', 12, False),
        ('Ora_Inizio_Turno', 18, False),
        ('Ora_Fine_Turno', 18, False),
        ('Ora_Inizio_Turno_Spezzato', 25, False),
        ('Ora_Fine_Turno_Spezzato', 25, False),

        # === STRAORDINARI (sfondo giallo) ===
        ('Inizio_Strao_1', 15, False),
        ('Fine_Strao_1', 15, False),
        ('Inizio_Strao_2', 15, False),
        ('Fine_Strao_2', 15, False),
        ('Inizio_Strao_3', 15, False),
        ('Fine_Strao_3', 15, False),

        # === PAUSE (sfondo viola) ===
        ('Inizio_Pausa_1', 15, False),
        ('Fine_Pausa_1', 15, False),
        ('Inizio_Pausa_2', 15, False),
        ('Fine_Pausa_2', 15, False),
        ('Inizio_Pausa_3', 15, False),
        ('Fine_Pausa_3', 15, False),
        ('Inizio_Pausa_4', 15, False),
        ('Fine_Pausa_4', 15, False),
        ('Inizio_Pausa_5', 15, False),
        ('Fine_Pausa_5', 15, False),

        # === GIUSTIFICATIVI (sfondo arancione) ===
        ('Tipo_Giust_1', 15, False),
        ('Inizio_Giust_1', 15, False),
        ('Fine_Giust_1', 15, False),
        ('Tipo_Giust_2', 15, False),
        ('Inizio_Giust_2', 15, False),
        ('Fine_Giust_2', 15, False),
        ('Tipo_Giust_3', 15, False),
        ('Inizio_Giust_3', 15, False),
        ('Fine_Giust_3', 15, False),
        ('Tipo_Giust_4', 15, False),
        ('Inizio_Giust_4', 15, False),
        ('Fine_Giust_4', 15, False),
        ('Tipo_Giust_5', 15, False),
        ('Inizio_Giust_5', 15, False),
        ('Fine_Giust_5', 15, False),

        # === SKILL E POSTAZIONE (sfondo grigio) ===
        ('Etichetta_Skill', 20, False),
        ('Postazione', 15, False),
    ]

    # Colori per sezioni
    colors = {
        'obbligatorio': 'FFCCCC',  # Rosso chiaro
        'anagrafica': 'CCE5FF',    # Azzurro
        'turno': 'CCFFCC',         # Verde chiaro
        'straordinari': 'FFFFCC',  # Giallo
        'pause': 'E5CCFF',         # Viola chiaro
        'giustificativi': 'FFDDAA', # Arancione
        'skill': 'E0E0E0'          # Grigio
    }

    # Determina colore per ogni colonna
    def get_color(idx):
        if idx < 4:
            return colors['obbligatorio']
        elif idx < 7:
            return colors['anagrafica']
        elif idx < 12:
            return colors['turno']
        elif idx < 18:
            return colors['straordinari']
        elif idx < 28:
            return colors['pause']
        elif idx < 43:
            return colors['giustificativi']
        else:
            return colors['skill']

    # Scrivi header
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for idx, (col_name, width, required) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=idx)
        cell.value = col_name
        cell.font = Font(bold=True, color='000000')
        cell.fill = PatternFill(start_color=get_color(idx-1), end_color=get_color(idx-1), fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

        # Imposta larghezza colonna
        ws.column_dimensions[get_column_letter(idx)].width = width

    # Aggiungi riga di esempio
    example_row = [
        'Mario',                # Nome
        'Rossi',                # Cognome
        '12345',                # ID_SAP
        '2025-01-15',           # Data_Riferimento
        'Full Time',            # Tipo_Contratto
        '1',                    # FTE
        '40',                   # Ore_Settimana
        'T1',                   # ID_Turno
        '09:00',                # Ora_Inizio_Turno
        '18:00',                # Ora_Fine_Turno
        '',                     # Ora_Inizio_Turno_Spezzato
        '',                     # Ora_Fine_Turno_Spezzato
        '18:00',                # Inizio_Strao_1
        '20:00',                # Fine_Strao_1
        '',                     # Inizio_Strao_2
        '',                     # Fine_Strao_2
        '',                     # Inizio_Strao_3
        '',                     # Fine_Strao_3
        '12:00',                # Inizio_Pausa_1
        '13:00',                # Fine_Pausa_1
        '',                     # Inizio_Pausa_2
        '',                     # Fine_Pausa_2
        '',                     # Inizio_Pausa_3
        '',                     # Fine_Pausa_3
        '',                     # Inizio_Pausa_4
        '',                     # Fine_Pausa_4
        '',                     # Inizio_Pausa_5
        '',                     # Fine_Pausa_5
        '',                     # Tipo_Giust_1
        '',                     # Inizio_Giust_1
        '',                     # Fine_Giust_1
        '',                     # Tipo_Giust_2
        '',                     # Inizio_Giust_2
        '',                     # Fine_Giust_2
        '',                     # Tipo_Giust_3
        '',                     # Inizio_Giust_3
        '',                     # Fine_Giust_3
        '',                     # Tipo_Giust_4
        '',                     # Inizio_Giust_4
        '',                     # Fine_Giust_4
        '',                     # Tipo_Giust_5
        '',                     # Inizio_Giust_5
        '',                     # Fine_Giust_5
        'Customer Care',        # Etichetta_Skill
        'Sede',                 # Postazione
    ]

    for idx, value in enumerate(example_row, 1):
        cell = ws.cell(row=2, column=idx)
        cell.value = value
        cell.border = border
        cell.alignment = Alignment(horizontal='left', vertical='center')

    # Congela prima riga
    ws.freeze_panes = 'A2'

    # Crea foglio istruzioni
    ws_instructions = wb.create_sheet("Istruzioni")

    instructions = [
        ["TEMPLATE IMPORT OPERATORI - ISTRUZIONI"],
        [""],
        ["COLONNE OBBLIGATORIE (sfondo rosso):"],
        ["  • Nome: Nome operatore"],
        ["  • Cognome: Cognome operatore"],
        ["  • ID_SAP: Codice univoco operatore (testo, es: '12345')"],
        ["  • Data_Riferimento: Data nel formato YYYY-MM-DD (es: 2025-01-15)"],
        [""],
        ["COLONNE OPZIONALI:"],
        [""],
        ["ANAGRAFICA (sfondo azzurro):"],
        ["  • Tipo_Contratto: Es. Full Time, Part Time, Apprendista"],
        ["  • FTE: Full Time Equivalent (es: 1 = 100%, 0.5 = 50%)"],
        ["  • Ore_Settimana: Ore contrattuali settimanali (es: 40)"],
        [""],
        ["TURNO (sfondo verde):"],
        ["  • ID_Turno: Identificativo turno (es: T1, MATTINA, POME)"],
        ["  • Ora_Inizio_Turno: Formato HH:MM (es: 09:00)"],
        ["  • Ora_Fine_Turno: Formato HH:MM (es: 18:00)"],
        ["  • Ora_Inizio_Turno_Spezzato: Se il turno è spezzato"],
        ["  • Ora_Fine_Turno_Spezzato: Se il turno è spezzato"],
        [""],
        ["STRAORDINARI (sfondo giallo): 3 slot disponibili"],
        ["  • Inizio_Strao_1/2/3: Inizio straordinario (HH:MM)"],
        ["  • Fine_Strao_1/2/3: Fine straordinario (HH:MM)"],
        [""],
        ["PAUSE (sfondo viola): 5 slot disponibili"],
        ["  • Inizio_Pausa_1/2/3/4/5: Inizio pausa (HH:MM)"],
        ["  • Fine_Pausa_1/2/3/4/5: Fine pausa (HH:MM)"],
        [""],
        ["GIUSTIFICATIVI (sfondo arancione): 5 slot disponibili"],
        ["  • Tipo_Giust_1/2/3/4/5: Assenza, Ferie, Malattia, Permesso, ROL, Congedo"],
        ["  • Inizio_Giust_1/2/3/4/5: Inizio assenza (HH:MM)"],
        ["  • Fine_Giust_1/2/3/4/5: Fine assenza (HH:MM)"],
        [""],
        ["SKILL E POSTAZIONE (sfondo grigio):"],
        ["  • Etichetta_Skill: Competenza principale (es: Customer Care, Tech Support)"],
        ["  • Postazione: Sede, Smart Working, Trasferta, Permesso, Assente"],
        [""],
        ["FORMATI SUPPORTATI:"],
        ["  • Date: YYYY-MM-DD, DD/MM/YYYY, MM/DD/YYYY"],
        ["  • Orari: HH:MM, HH:MM:SS"],
        [""],
        ["NOTE:"],
        ["  • Lascia le celle vuote per campi opzionali non necessari"],
        ["  • Se ID_SAP + Data_Riferimento esistono già, il record viene AGGIORNATO"],
        ["  • La prima riga (ESEMPIO) può essere eliminata o modificata"],
        ["  • Puoi aggiungere quante righe vuoi sotto l'esempio"],
        [""],
        ["IMPORT DA SCRIPT:"],
        ["  python scripts/import_excel_operatori.py --file percorso/file.xlsx"],
        [""],
        ["IMPORT DA GUI:"],
        ["  Tab Anagrafica → Importa Excel → Seleziona questo file"],
    ]

    for row_idx, row_data in enumerate(instructions, 1):
        cell = ws_instructions.cell(row=row_idx, column=1)
        cell.value = row_data[0]

        # Grassetto per titoli
        if row_data[0].isupper() or row_data[0].endswith(':'):
            cell.font = Font(bold=True)

    ws_instructions.column_dimensions['A'].width = 80

    # Salva
    import os
    os.makedirs('templates', exist_ok=True)
    wb.save(output_path)

    print(f"✓ Template creato: {output_path}")
    print(f"  Foglio 'Operatori': Contiene header + esempio")
    print(f"  Foglio 'Istruzioni': Guida completa")


if __name__ == '__main__':
    create_template()
