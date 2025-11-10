"""
Script per creare template Excel per import operatori massivo

Genera un file Excel con:
- Tutte le colonne necessarie per import completo
- Esempi precompilati
- Istruzioni dettagliate
"""

import pandas as pd
import sys
import os
from datetime import datetime

def crea_template_operatori(output_file='template_import_operatori.xlsx'):
    """Crea template Excel per import operatori"""

    print("\n=== CREAZIONE TEMPLATE EXCEL OPERATORI ===\n")

    # Dati di esempio (3 operatori)
    data = {
        # === CAMPI OBBLIGATORI ===
        'Nome': ['Mario', 'Lucia', 'Giovanni', ''],
        'Cognome': ['Rossi', 'Bianchi', 'Verdi', ''],
        'ID_SAP': ['SAP001', 'SAP002', 'SAP003', ''],
        'Data_Riferimento': ['2025-11-08', '2025-11-08', '2025-11-08', ''],

        # === DATI CONTRATTUALI ===
        'Tipo_Contratto': ['Full Time', 'Part Time', 'Full Time', ''],
        'FTE': [1.0, 0.5, 1.0, ''],
        'Ore_Settimana': [40, 20, 40, ''],

        # === TURNO ORDINARIO ===
        'ID_Turno': ['T1', 'T2', 'T1', ''],
        'Ora_Inizio_Turno': ['09:00', '14:00', '09:00', ''],
        'Ora_Fine_Turno': ['18:00', '22:00', '18:00', ''],

        # === TURNO SPEZZATO (opzionale) ===
        'Ora_Inizio_Turno_Spezzato': ['', '', '', ''],
        'Ora_Fine_Turno_Spezzato': ['', '', '', ''],

        # === STRAORDINARI (3 slot) ===
        'Inizio_Strao_1': ['18:00', '', '', ''],
        'Fine_Strao_1': ['20:00', '', '', ''],
        'Inizio_Strao_2': ['', '', '', ''],
        'Fine_Strao_2': ['', '', '', ''],
        'Inizio_Strao_3': ['', '', '', ''],
        'Fine_Strao_3': ['', '', '', ''],

        # === PAUSE (5 slot) ===
        'Inizio_Pausa_1': ['10:45', '16:30', '10:45', ''],
        'Fine_Pausa_1': ['11:00', '16:45', '11:00', ''],
        'Inizio_Pausa_2': ['13:00', '', '13:00', ''],
        'Fine_Pausa_2': ['14:00', '', '14:00', ''],
        'Inizio_Pausa_3': ['', '', '', ''],
        'Fine_Pausa_3': ['', '', '', ''],
        'Inizio_Pausa_4': ['', '', '', ''],
        'Fine_Pausa_4': ['', '', '', ''],
        'Inizio_Pausa_5': ['', '', '', ''],
        'Fine_Pausa_5': ['', '', '', ''],

        # === GIUSTIFICATIVI (5 slot) ===
        'Tipo_Giust_1': ['', 'Permesso', '', ''],
        'Inizio_Giust_1': ['', '20:00', '', ''],
        'Fine_Giust_1': ['', '22:00', '', ''],
        'Tipo_Giust_2': ['', '', '', ''],
        'Inizio_Giust_2': ['', '', '', ''],
        'Fine_Giust_2': ['', '', '', ''],
        'Tipo_Giust_3': ['', '', '', ''],
        'Inizio_Giust_3': ['', '', '', ''],
        'Fine_Giust_3': ['', '', '', ''],
        'Tipo_Giust_4': ['', '', '', ''],
        'Inizio_Giust_4': ['', '', '', ''],
        'Fine_Giust_4': ['', '', '', ''],
        'Tipo_Giust_5': ['', '', '', ''],
        'Inizio_Giust_5': ['', '', '', ''],
        'Fine_Giust_5': ['', '', '', ''],

        # === SKILL E POSTAZIONE ===
        'Etichetta_Skill': ['CUSTOMER_CARE', 'BACK_OFFICE', 'TECHNICAL_SUPPORT', ''],
        'Microskill': ['Premium', '', 'Advanced', ''],
        'Postazione': ['Sede', 'Smart Working', 'Sede', ''],
    }

    # Crea DataFrame
    df = pd.DataFrame(data)

    # Crea Excel con formattazione
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        # Scrivi dati
        df.to_excel(writer, sheet_name='Operatori', index=False)

        # Ottieni worksheet
        worksheet = writer.sheets['Operatori']

        # Imposta larghezze colonne
        worksheet.column_dimensions['A'].width = 12  # Nome
        worksheet.column_dimensions['B'].width = 12  # Cognome
        worksheet.column_dimensions['C'].width = 12  # ID_SAP
        worksheet.column_dimensions['D'].width = 15  # Data_Riferimento
        worksheet.column_dimensions['E'].width = 15  # Tipo_Contratto
        worksheet.column_dimensions['F'].width = 8   # FTE
        worksheet.column_dimensions['G'].width = 13  # Ore_Settimana

        # Formatta header
        from openpyxl.styles import Font, PatternFill, Alignment

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=10, color="FFFFFF")

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # Colora sezioni diverse
        # Dati contrattuali (E-G) - Verde
        contract_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        for col in ['E', 'F', 'G']:
            worksheet[f'{col}1'].fill = contract_fill

        # Turno (H-L) - Giallo
        turno_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        for col in ['H', 'I', 'J', 'K', 'L']:
            worksheet[f'{col}1'].fill = turno_fill

        # Straordinari (M-R) - Arancione
        strao_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        for col in ['M', 'N', 'O', 'P', 'Q', 'R']:
            worksheet[f'{col}1'].fill = strao_fill

        # === SHEET ISTRUZIONI ===
        istruzioni = pd.DataFrame({
            'ISTRUZIONI PER L\'IMPORT MASSIVO OPERATORI': [
                '',
                '=== GUIDA RAPIDA ===',
                '',
                'Questo template permette di importare TUTTI i dati degli operatori in un colpo solo:',
                '✓ Dati anagrafici (Nome, Cognome, ID SAP)',
                '✓ Dati contrattuali (Tipo contratto, FTE, Ore settimanali)',
                '✓ Turni (ordinari e spezzati)',
                '✓ Straordinari (fino a 3 slot)',
                '✓ Pause (fino a 5 slot)',
                '✓ Giustificativi (fino a 5 slot con tipo e orari)',
                '✓ Skill e Postazione',
                '',
                '=== COLONNE OBBLIGATORIE ===',
                '',
                '1. Nome (testo)',
                '2. Cognome (testo)',
                '3. ID_SAP (codice univoco operatore)',
                '4. Data_Riferimento (formato: YYYY-MM-DD es: 2025-11-08)',
                '',
                'IMPORTANTE: Se manca anche solo UNA di queste colonne, la riga verrà saltata!',
                '',
                '=== COLONNE OPZIONALI ===',
                '',
                'DATI CONTRATTUALI:',
                '• Tipo_Contratto: Full Time, Part Time, Tempo Determinato',
                '• FTE: Numero decimale (es: 1.0 = full time, 0.5 = part time)',
                '• Ore_Settimana: Numero ore settimanali (es: 40, 20)',
                '',
                'TURNO:',
                '• ID_Turno: Codice turno (es: T1, T2) - deve esistere in tabella Turni',
                '• Ora_Inizio_Turno: Formato HH:MM (es: 09:00)',
                '• Ora_Fine_Turno: Formato HH:MM (es: 18:00)',
                '• Ora_Inizio_Turno_Spezzato: Inizio parte spezzata (es: 17:00)',
                '• Ora_Fine_Turno_Spezzato: Fine parte spezzata (es: 21:00)',
                '',
                'STRAORDINARI (3 slot):',
                '• Inizio_Strao_1, Fine_Strao_1',
                '• Inizio_Strao_2, Fine_Strao_2',
                '• Inizio_Strao_3, Fine_Strao_3',
                '  Formato sempre HH:MM (es: 18:00, 20:00)',
                '',
                'PAUSE (5 slot):',
                '• Inizio_Pausa_1, Fine_Pausa_1',
                '• Inizio_Pausa_2, Fine_Pausa_2',
                '• ... fino a Pausa_5',
                '  Formato sempre HH:MM (es: 10:45, 11:00)',
                '',
                'GIUSTIFICATIVI (5 slot):',
                '• Tipo_Giust_1: Assenza, Ferie, Malattia, Permesso, ROL, Congedo',
                '• Inizio_Giust_1, Fine_Giust_1: Orari formato HH:MM',
                '• ... fino a Giust_5',
                '',
                'SKILL E POSTAZIONE:',
                '• Etichetta_Skill: CUSTOMER_CARE, BACK_OFFICE, TECHNICAL_SUPPORT, ecc.',
                '• Microskill: Sottocategoria skill (es: Premium, Basic, Advanced) - opzionale',
                '  Il microskill è trasversale: un operatore con microskill "Premium"',
                '  può lavorare su qualsiasi skill che richiede "Premium"',
                '• Postazione: Sede, Smart Working, Trasferta, Permesso, Assente',
                '',
                '=== ESEMPI FORNITI ===',
                '',
                'Nel foglio "Operatori" trovi 3 esempi completi:',
                '',
                '1. Mario Rossi (SAP001):',
                '   - Turno 09:00-18:00',
                '   - Straordinari 18:00-20:00',
                '   - 2 pause (10:45-11:00, 13:00-14:00)',
                '   - Skill: CUSTOMER_CARE',
                '   - Microskill: Premium',
                '   - Postazione: Sede',
                '',
                '2. Lucia Bianchi (SAP002):',
                '   - Part time 0.5 FTE, 20h/settimana',
                '   - Turno 14:00-22:00',
                '   - 1 pausa (16:30-16:45)',
                '   - Giustificativo: Permesso 20:00-22:00',
                '   - Postazione: Smart Working',
                '',
                '3. Giovanni Verdi (SAP003):',
                '   - Turno standard 09:00-18:00',
                '   - 2 pause',
                '   - Skill: TECHNICAL_SUPPORT',
                '   - Microskill: Advanced',
                '',
                '=== COME USARE IL TEMPLATE ===',
                '',
                'PASSO 1: Compila il foglio "Operatori"',
                '  - Modifica gli esempi o aggiungine di nuovi',
                '  - Puoi eliminare le righe di esempio se non servono',
                '  - Aggiungi tante righe quanti sono i tuoi operatori',
                '',
                'PASSO 2: Verifica i dati',
                '  - Controlla che tutti gli ID_SAP siano univoci',
                '  - Verifica formato orari (sempre HH:MM)',
                '  - Verifica formato date (sempre YYYY-MM-DD)',
                '',
                'PASSO 3: Salva il file Excel',
                '',
                'PASSO 4: Import nel database',
                '  Apri il Prompt dei comandi (CMD) e digita:',
                '',
                '  cd C:\\Matrici',
                '  python scripts/import_excel_operatori.py --file template_import_operatori.xlsx',
                '',
                '  Oppure se usi un nome diverso:',
                '  python scripts/import_excel_operatori.py --file mio_file.xlsx',
                '',
                '  Per specificare un foglio diverso:',
                '  python scripts/import_excel_operatori.py --file mio_file.xlsx --sheet "MioFoglio"',
                '',
                '=== COSA FA LO SCRIPT DI IMPORT ===',
                '',
                '✓ Legge il file Excel',
                '✓ Valida i dati (campi obbligatori, formati)',
                '✓ Controlla se operatori esistono già (stesso ID_SAP e Data)',
                '✓ Se esistono → AGGIORNA i dati',
                '✓ Se non esistono → INSERISCE nuovi operatori',
                '✓ Mostra report dettagliato:',
                '  - Quanti operatori importati',
                '  - Quanti aggiornati',
                '  - Eventuali errori',
                '',
                '=== NOTE IMPORTANTI ===',
                '',
                '⚠ DUPLICATI:',
                'Se hai già un operatore con stesso ID_SAP e Data_Riferimento,',
                'i suoi dati verranno SOVRASCRITTI con quelli del file Excel!',
                '',
                '⚠ FORMATI ORARI:',
                'Excel può cambiare i formati automaticamente.',
                'Assicurati che gli orari siano sempre in formato HH:MM (testo).',
                'Se Excel li converte in formato ora, lo script li gestisce automaticamente.',
                '',
                '⚠ ID_TURNO:',
                'Se specifichi un ID_Turno, deve esistere nella tabella Turni.',
                'Prima importa i turni con: python scripts/import_excel_turni.py',
                '',
                '⚠ DATA_RIFERIMENTO:',
                'Ogni operatore può avere dati diversi per date diverse.',
                'Esempio: Mario il 08/11 fa 09:00-18:00, il 09/11 fa 14:00-22:00',
                'Basta creare 2 righe con stesso ID_SAP ma date diverse.',
                '',
                '=== DOMANDE FREQUENTI ===',
                '',
                'Q: Posso importare operatori per più giorni contemporaneamente?',
                'A: Sì! Basta ripetere la riga con ID_SAP uguale ma Data_Riferimento diversa.',
                '',
                'Q: Cosa succede se lascio campi vuoti?',
                'A: I campi opzionali vuoti vengono salvati come NULL (nessun dato).',
                '',
                'Q: Posso reimportare lo stesso file più volte?',
                'A: Sì, gli operatori esistenti vengono aggiornati automaticamente.',
                '',
                'Q: Come elimino un operatore?',
                'A: Dall\'interfaccia Matrici, seleziona l\'operatore e clicca "Elimina".',
                '',
                'Q: Posso modificare manualmente dopo l\'import?',
                'A: Sì! Apri Matrici, fai doppio click sull\'operatore e modifica.',
                '',
                '=== SUPPORTO ===',
                '',
                'Per problemi o domande, controlla:',
                '• Le colonne nel foglio "Operatori"',
                '• I formati di data e ora',
                '• Gli errori mostrati dallo script di import',
                '',
            ]
        })

        istruzioni.to_excel(writer, sheet_name='Istruzioni', index=False, header=False)

        # Formatta foglio istruzioni
        ws_istr = writer.sheets['Istruzioni']
        ws_istr.column_dimensions['A'].width = 85

        # Grassetto per titoli
        from openpyxl.styles import Font
        bold_font = Font(bold=True, size=12)

        for row_idx in [2, 13, 20, 48, 66, 91, 105, 121]:
            ws_istr[f'A{row_idx}'].font = bold_font

    print(f"✅ Template creato: {output_file}")
    print(f"\nIl file contiene:")
    print(f"  - Sheet 'Operatori': 3 esempi completi + riga vuota")
    print(f"  - Sheet 'Istruzioni': Guida dettagliata")
    print(f"\nColonne totali: {len(data)}")
    print(f"  • 4 obbligatorie (Nome, Cognome, ID_SAP, Data)")
    print(f"  • {len(data) - 4} opzionali (turni, straordinari, pause, giustificativi, ecc.)")
    print(f"\nProssimi passi:")
    print(f"  1. Apri {output_file} con Excel")
    print(f"  2. Compila i tuoi operatori")
    print(f"  3. Salva il file")
    print(f"  4. Import: python scripts/import_excel_operatori.py --file {output_file}")
    print()


if __name__ == '__main__':
    # Verifica se è stato specificato un nome file
    if len(sys.argv) > 1:
        output_file = sys.argv[1]
    else:
        output_file = 'template_import_operatori.xlsx'

    # Crea template
    try:
        crea_template_operatori(output_file)
        print("✅ Template creato con successo!\n")
    except Exception as e:
        print(f"❌ Errore durante la creazione del template: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
