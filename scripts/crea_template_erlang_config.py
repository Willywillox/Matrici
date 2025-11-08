"""
Genera template Excel per configurazione parametri Erlang C

Il file Excel conterrà configurazioni per calcoli workforce planning:
- AHT (Average Handle Time)
- Shrinkage (pause, formazione, ecc.)
- Service Level targets
- Occupancy target
"""

import pandas as pd
import sys
from pathlib import Path


def crea_template_erlang(output_file='template_erlang_config.xlsx'):
    """
    Crea template Excel per import configurazione Erlang C

    Args:
        output_file: Nome file output
    """

    # Dati di esempio (Voice + Chat)
    config_data = {
        'Skill': [
            'CUSTOMER_CARE_VOICE',
            'CUSTOMER_CARE_CHAT',
            'BACK_OFFICE',
            'TECHNICAL_SUPPORT',
            'SALES_CHAT',
            ''
        ],
        'Tipo_Canale': [
            'Voice',
            'Chat',
            'Email',
            'Voice',
            'Chat',
            'Voice'
        ],
        'AHT_Seconds': [
            180,   # 3 minuti voice
            240,   # 4 minuti chat
            360,   # 6 minuti email
            300,   # 5 minuti
            180,   # 3 minuti chat
            180
        ],
        'Concurrency': [
            1,     # Voice: sempre 1
            3,     # Chat: 3 chat simultanee
            1,     # Email: 1
            1,     # Voice: 1
            4,     # Chat: 4 chat simultanee
            1
        ],
        'Tempo_Pausa_Minuti': [
            0,     # Pause gestite separatamente
            0,
            0,
            0,
            0,
            0
        ],
        'Shrinkage': [
            0.30,  # 30% (pause, formazione, riunioni, ecc.)
            0.28,  # 28% chat
            0.25,  # 25%
            0.30,  # 30%
            0.27,  # 27%
            0.30
        ],
        'Service_Level_Target': [
            0.80,  # 80% delle chiamate (per voice)
            0.80,  # Per chat usa ASA
            0.75,  # 75%
            0.80,  # 80%
            0.80,  # Per chat usa ASA
            0.80
        ],
        'Service_Level_Seconds': [
            20,    # Entro 20 secondi (voice)
            20,    # Ignorato per chat
            14400, # 4 ore per email
            20,    # Entro 20 secondi
            15,    # Ignorato per chat
            20
        ],
        'ASA_Target_Seconds': [
            0,     # Ignorato per voice
            60,    # 60s tempo prima risposta chat
            0,     # Ignorato per email
            0,     # Ignorato per voice
            45,    # 45s tempo prima risposta chat
            0
        ],
        'Occupancy_Target': [
            0.85,  # 85% occupancy
            0.80,  # 80% chat
            0.80,  # 80%
            0.85,  # 85%
            0.82,  # 82% chat
            0.85
        ],
        'Interval_Minutes': [
            30,    # Calcoli ogni 30 minuti
            30,
            30,
            30,
            30,
            30
        ],
        'Note': [
            'Assistenza clienti telefono',
            'Assistenza clienti chat con 3 chat simultanee',
            'Attività back office email',
            'Supporto tecnico avanzato',
            'Vendite chat con 4 chat simultanee',
            ''
        ]
    }

    df = pd.DataFrame(config_data)

    # Crea Excel con formattazione
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Erlang_Config', index=False)

        # Ottieni workbook e worksheet
        workbook = writer.book
        worksheet = writer.sheets['Erlang_Config']

        # Formattazione header
        from openpyxl.styles import Font, PatternFill, Alignment

        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Larghezza colonne
        column_widths = {
            'A': 28,  # Skill
            'B': 15,  # Tipo_Canale
            'C': 18,  # AHT_Seconds
            'D': 15,  # Concurrency
            'E': 22,  # Tempo_Pausa_Minuti
            'F': 15,  # Shrinkage
            'G': 22,  # Service_Level_Target
            'H': 25,  # Service_Level_Seconds
            'I': 20,  # ASA_Target_Seconds
            'J': 20,  # Occupancy_Target
            'K': 20,  # Interval_Minutes
            'L': 50   # Note
        }

        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width

        # Aggiungi sheet con istruzioni
        istruzioni = pd.DataFrame({
            'ISTRUZIONI PER LA CONFIGURAZIONE ERLANG C': [
                '',
                '=== PARAMETRI ===',
                '',
                'Skill: Codice skill/coda (deve corrispondere a Skills configurati)',
                '',
                'Tipo_Canale: Tipo di canale (Voice/Chat/Email)',
                '  - Voice: Telefonia tradizionale, concurrency = 1',
                '  - Chat: Chat testuale, concurrency > 1 (chat simultanee)',
                '  - Email: Email, concurrency = 1',
                '',
                'AHT_Seconds: Average Handle Time in secondi',
                '  - Tempo medio gestione chiamata/contatto',
                '  - Include tempo conversazione + after call work',
                '  - Esempio Voice: 180 = 3 minuti',
                '  - Esempio Chat: 240 = 4 minuti',
                '',
                'Concurrency: Chat simultanee per agente',
                '  - Voice/Email: sempre 1',
                '  - Chat: tipicamente 2-5 (quante chat gestite contemporaneamente)',
                '  - Esempio: 3 = agente gestisce 3 chat in parallelo',
                '',
                'Tempo_Pausa_Minuti: Minuti pausa per ora (normalmente 0)',
                '  - Pause gestite separatamente nello shrinkage',
                '  - Lasciare a 0 se pause già incluse in Shrinkage',
                '',
                'Shrinkage: Percentuale tempo non produttivo (0-1)',
                '  - Include: pause, formazione, riunioni, assenze',
                '  - Esempio: 0.30 = 30% del tempo non produttivo',
                '  - Tipicamente tra 25% e 35%',
                '',
                'Service_Level_Target: Target Service Level (0-1) [per Voice/Email]',
                '  - Percentuale chiamate da rispondere entro target',
                '  - Esempio: 0.80 = 80% delle chiamate',
                '  - Standard Voice: 0.80 (80/20)',
                '  - Per Chat: questo parametro viene ignorato, usa ASA',
                '',
                'Service_Level_Seconds: Secondi target risposta [per Voice/Email]',
                '  - Tempo massimo attesa per Service Level',
                '  - Voice: tipicamente 20 secondi',
                '  - Email: tipicamente 14400 (4 ore)',
                '  - Per Chat: ignorato',
                '',
                'ASA_Target_Seconds: Average Speed to Answer target [per Chat]',
                '  - Tempo medio entro cui rispondere alla prima risposta',
                '  - Solo per Chat, esempio: 60 = rispondere entro 1 minuto mediamente',
                '  - Tipicamente 30-90 secondi per chat',
                '  - Per Voice/Email: lasciare a 0',
                '',
                'Occupancy_Target: Target occupancy agenti (0-1)',
                '  - Percentuale tempo in chiamata/lavoro',
                '  - Esempio: 0.85 = 85% occupati',
                '  - Voice: tipicamente 80-90%',
                '  - Chat: tipicamente 75-85% (più basso per gestire simultaneità)',
                '',
                'Interval_Minutes: Intervallo calcolo in minuti',
                '  - Tipicamente 30 minuti',
                '  - Deve corrispondere a intervalli forecast',
                '',
                'Note: Descrizione/annotazioni',
                '',
                '=== ESEMPI VALORI TIPICI ===',
                '',
                'Voice - Customer Care:',
                '  - Tipo_Canale: Voice',
                '  - AHT: 180-240 secondi (3-4 minuti)',
                '  - Concurrency: 1 (sempre)',
                '  - Shrinkage: 30%',
                '  - Service Level: 80% in 20 secondi',
                '  - Occupancy: 85%',
                '',
                'Chat - Customer Care:',
                '  - Tipo_Canale: Chat',
                '  - AHT: 180-300 secondi (3-5 minuti per chat)',
                '  - Concurrency: 3-4 (chat simultanee)',
                '  - Shrinkage: 28%',
                '  - ASA Target: 60 secondi (tempo prima risposta)',
                '  - Occupancy: 80%',
                '',
                'Email - Back Office:',
                '  - Tipo_Canale: Email',
                '  - AHT: 300-600 secondi (5-10 minuti)',
                '  - Concurrency: 1',
                '  - Shrinkage: 25%',
                '  - Service Level: 75% in 4 ore (14400s)',
                '  - Occupancy: 80%',
                '',
                'Voice - Technical Support:',
                '  - Tipo_Canale: Voice',
                '  - AHT: 300-600 secondi (5-10 minuti)',
                '  - Concurrency: 1',
                '  - Shrinkage: 30%',
                '  - Service Level: 80% in 20 secondi',
                '  - Occupancy: 85%',
                '',
                '=== IMPORT ===',
                '',
                'Dopo aver completato la configurazione:',
                '  python scripts/import_erlang_config.py template_erlang_config.xlsx',
                ''
            ]
        })

        istruzioni.to_excel(writer, sheet_name='Istruzioni', index=False, header=False)
        worksheet_istr = writer.sheets['Istruzioni']
        worksheet_istr.column_dimensions['A'].width = 80

        # Formattazione istruzioni
        for row in worksheet_istr.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if cell.value.startswith('==='):
                        cell.font = Font(bold=True, size=12, color='1F4E78')
                    elif cell.value.startswith('  '):
                        cell.font = Font(italic=True, size=10)

    print(f"\n✅ Template creato: {output_file}")
    print("\nColonne create:")
    print("  - Skill: Codice skill/coda")
    print("  - Tipo_Canale: Voice/Chat/Email")
    print("  - AHT_Seconds: Average Handle Time")
    print("  - Concurrency: Chat simultanee per agente (1 per Voice)")
    print("  - Tempo_Pausa_Minuti: Minuti pausa/ora")
    print("  - Shrinkage: % tempo non produttivo")
    print("  - Service_Level_Target: Target % SL (per Voice/Email)")
    print("  - Service_Level_Seconds: Secondi target risposta (per Voice/Email)")
    print("  - ASA_Target_Seconds: Tempo prima risposta (per Chat)")
    print("  - Occupancy_Target: Target occupancy")
    print("  - Interval_Minutes: Intervallo calcolo")
    print("  - Note: Annotazioni")
    print("\n5 configurazioni di esempio incluse (Voice + Chat + Email)")
    print("Sheet 'Istruzioni' con guida completa")
    print(f"\nPer importare: python scripts/import_erlang_config.py {output_file}\n")


if __name__ == '__main__':
    output = sys.argv[1] if len(sys.argv) > 1 else 'template_erlang_config.xlsx'
    crea_template_erlang(output)
